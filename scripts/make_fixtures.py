#!/usr/bin/env python3
"""Build deterministic synthetic DOCX occurrence fixtures with bundled Python.

Fixture authoring uses python-docx; verification in tests/oracle.py uses only
zipfile and ElementTree and never imports the application or this generator.
"""
from copy import deepcopy
from io import BytesIO
import json
from pathlib import Path
import struct
import zipfile
import zlib

from docx import Document
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "fixtures"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
      "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
      "pic": "http://schemas.openxmlformats.org/drawingml/2006/picture",
      "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
      "pr": "http://schemas.openxmlformats.org/package/2006/relationships",
      "ct": "http://schemas.openxmlformats.org/package/2006/content-types",
      "mc": "http://schemas.openxmlformats.org/markup-compatibility/2006",
      "v": "urn:schemas-microsoft-com:vml",
      "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
      "dgm": "http://schemas.openxmlformats.org/drawingml/2006/diagram",
      "wpg": "http://schemas.microsoft.com/office/word/2010/wordprocessingGroup",
      "adec": "http://schemas.microsoft.com/office/drawing/2017/decorative"}


def tag(prefix, local):
    return "{" + NS[prefix] + "}" + local


def png_icon():
    """A literal synthetic 48 px diamond; no third-party image or private data."""
    rows = []
    for y in range(48):
        row = bytearray([0])
        for x in range(48):
            color = (24, 74, 65) if abs(x - 23.5) + abs(y - 23.5) < 16 else (216, 237, 226)
            row.extend(color)
        rows.append(bytes(row))
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data) & 0xFFFFFFFF)
    return b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", 48, 48, 8, 2, 0, 0, 0)) + chunk(b"IDAT", zlib.compress(b"".join(rows))) + chunk(b"IEND", b"")


def document(title, explanation):
    doc = Document()
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        for border in list(style.element.iter(tag("w", "pBdr"))):
            border.getparent().remove(border)
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_paragraph(title, "Title")
    doc.add_paragraph(explanation)
    return doc


def picture(paragraph, descr=None, title=None, name=None):
    drawing = paragraph.add_run().add_picture(BytesIO(png_icon()), width=Inches(.38))
    properties = drawing._inline.docPr
    if descr is not None:
        properties.set("descr", descr)
    if title is not None:
        properties.set("title", title)
    if name:
        properties.set("name", name)
    return drawing._inline


def floating(inline):
    """Convert an ordinary inline to an ordinary anchored image, retaining size."""
    anchor = etree.Element(tag("wp", "anchor"), nsmap=inline.nsmap,
                           distT="0", distB="0", distL="114300", distR="114300",
                           simplePos="0", relativeHeight="0", behindDoc="0", locked="0",
                           layoutInCell="1", allowOverlap="1")
    etree.SubElement(anchor, tag("wp", "simplePos"), x="0", y="0")
    horizontal = etree.SubElement(anchor, tag("wp", "positionH"), relativeFrom="column")
    etree.SubElement(horizontal, tag("wp", "posOffset")).text = "4200000"
    vertical = etree.SubElement(anchor, tag("wp", "positionV"), relativeFrom="paragraph")
    etree.SubElement(vertical, tag("wp", "posOffset")).text = "0"
    children = list(inline)
    for child in children:
        if child.tag in (tag("wp", "extent"), tag("wp", "effectExtent")):
            anchor.append(child)
    etree.SubElement(anchor, tag("wp", "wrapNone"))
    for child in children:
        if child.getparent() is inline:
            anchor.append(child)
    inline.getparent().replace(inline, anchor)
    return anchor


def packed(doc):
    data = BytesIO()
    doc.save(data)
    with zipfile.ZipFile(BytesIO(data.getvalue())) as archive:
        return {entry.filename: archive.read(entry.filename) for entry in archive.infolist()}


def xml_bytes(root):
    return etree.tostring(root, encoding="UTF-8", xml_declaration=True, standalone=True)


def add_member(members, name, data, content_type):
    members[name] = data
    types = etree.fromstring(members["[Content_Types].xml"])
    etree.SubElement(types, tag("ct", "Override"), PartName="/" + name, ContentType=content_type)
    members["[Content_Types].xml"] = xml_bytes(types)


def save(members, filename):
    OUT.mkdir(parents=True, exist_ok=True)
    path = OUT / filename
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as archive:
        for name in sorted(members):
            info = zipfile.ZipInfo(name, date_time=(2025, 1, 1, 0, 0, 0))
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o600 << 16
            archive.writestr(info, members[name])
    return path


def make_bench():
    doc = document("Image review bench", "Review each image occurrence in its own context. The repeated status symbols share identical image bytes and original descriptions.")
    header = doc.sections[0].header.paragraphs[0]
    header.add_run("Shared report header   ")
    picture(header, "Header symbol", name="Shared header icon")
    footer = doc.sections[0].footer.paragraphs[0]
    footer.add_run("Internal review copy   ")
    picture(footer, "", "Footer badge", "Footer icon")
    for heading, label in (("Intake status", "Intake icon"), ("Release status", "Release icon")):
        doc.add_paragraph(heading, "Heading 1")
        p = doc.add_paragraph("Status checkpoint   ")
        picture(p, "Status icon", "Status", label)
    doc.add_paragraph("Summary marker", "Heading 1")
    p = doc.add_paragraph("The floating marker is anchored to this summary paragraph.")
    floating(picture(p, None, "Floating marker", "Floating summary icon"))
    doc.add_section(WD_SECTION.NEW_PAGE)
    doc.add_paragraph("Shared report context", "Heading 1")
    doc.add_paragraph("This second section deliberately reuses the same header and footer parts. A shared header image is one stored occurrence, even when Word displays it on multiple pages.")
    # Explicit references in both section property elements prove part deduplication.
    sections = doc._element.xpath(".//w:sectPr")
    for index, kind in enumerate(("headerReference", "footerReference")):
        first = sections[0].find(tag("w", kind))
        assert first is not None
        if sections[-1].find(tag("w", kind)) is None:
            sections[-1].insert(index, deepcopy(first))
    members = packed(doc)
    # Formatting witnesses keep the same XML information set: literal tab
    # normalization, single quotes, character-reference space, and an explicit
    # start/end docPr pair rather than a self-closing element.
    members["word/document.xml"] = members["word/document.xml"].replace(
        b'descr="Status icon"', b"descr='Status\ticon'", 1)
    members["word/document.xml"] = members["word/document.xml"].replace(
        b'<wp:docPr id="3" name="Floating summary icon" title="Floating marker"/>',
        b'<wp:docPr id="3" name="Floating summary icon" title="Floating marker"></wp:docPr>', 1)
    members["word/header1.xml"] = members["word/header1.xml"].replace(
        b'descr="Header symbol"', b'descr="Header&#x20;symbol"', 1)
    add_member(members, "customXml/witness.bin", b"ALTSLOT UNKNOWN MEMBER\x00\xff\x01\nPreserve these bytes exactly.\n", "application/octet-stream")
    save(members, "bench.docx")
    changed = dict(members)
    changed["word/document.xml"] = changed["word/document.xml"].replace(b"Image review bench", b"Changed source bench", 1)
    save(changed, "bench-changed.docx")


def make_exclusions():
    doc = document("Unsupported image review cases", "These synthetic objects exercise exclusions. Only the plain control image is an editable supported occurrence.")
    doc.add_paragraph("Plain control", "Heading 1")
    picture(doc.add_paragraph("Editable control   "), "Control image", "Control", "Control")
    variants = {}
    for label in ("Tracked", "Chart", "SmartArt", "Grouped", "External", "AlternateContent", "Decorative", "Conflicting"):
        doc.add_paragraph(label, "Heading 2")
        variants[label] = picture(doc.add_paragraph(label + " object   "), label + " image", label, label)
    # Tracked changes enclose the drawing run, never the whole paragraph.
    inline = variants["Tracked"]
    run = inline.getparent().getparent()
    parent = run.getparent()
    insertion = etree.Element(tag("w", "ins"), {tag("w", "id"): "1", tag("w", "author"): "Synthetic reviewer", tag("w", "date"): "2025-01-01T00:00:00Z"})
    parent.replace(run, insertion)
    insertion.append(run)
    for label, prefix, local, uri in (("Chart", "c", "chart", NS["c"]), ("SmartArt", "dgm", "relIds", NS["dgm"])):
        graphic = variants[label].find(".//" + tag("a", "graphicData"))
        graphic.clear()
        graphic.set("uri", uri)
        child = etree.SubElement(graphic, tag(prefix, local))
        child.set(tag("r", "id" if label == "Chart" else "dm"), "rIdFixture" + label)
    graphic = variants["Grouped"].find(".//" + tag("a", "graphicData"))
    pic = graphic.find(tag("pic", "pic"))
    graphic.remove(pic)
    graphic.set("uri", NS["wpg"])
    group = etree.SubElement(graphic, tag("wpg", "wgp"))
    etree.SubElement(group, tag("wpg", "cNvGrpSpPr"))
    etree.SubElement(group, tag("wpg", "grpSpPr"))
    group.append(pic)
    blip = variants["External"].find(".//" + tag("a", "blip"))
    blip.attrib.pop(tag("r", "embed"))
    blip.set(tag("r", "link"), "rIdFixtureExternal")
    inline = variants["AlternateContent"]
    drawing = inline.getparent()
    run = drawing.getparent()
    alternate = etree.Element(tag("mc", "AlternateContent"), nsmap={"mc": NS["mc"], "wpg": NS["wpg"]})
    choice = etree.SubElement(alternate, tag("mc", "Choice"), Requires="wpg")
    run.replace(drawing, alternate)
    choice.append(drawing)
    etree.SubElement(alternate, tag("mc", "Fallback"))
    properties = variants["Decorative"].find(tag("wp", "docPr"))
    extension_list = etree.SubElement(properties, tag("a", "extLst"))
    extension = etree.SubElement(extension_list, tag("a", "ext"), uri="{C183D7F6-B498-43B3-948B-1728B52AA6E4}")
    etree.SubElement(extension, tag("adec", "decorative"), val="1")
    conflict = variants["Conflicting"].find(".//" + tag("pic", "cNvPr"))
    conflict.set("descr", "Different carrier description")
    conflict.set("title", "Different carrier title")
    doc.add_paragraph("Legacy VML", "Heading 2")
    p = doc.add_paragraph("Legacy VML object   ")
    run = p.add_run()._r
    pict = etree.SubElement(run, tag("w", "pict"))
    shape = etree.SubElement(pict, tag("v", "shape"), id="FixtureVML", style="width:28pt;height:28pt", type="#_x0000_t75")
    image_rel_id = next(key for key, rel in doc.part.rels.items() if rel.reltype.endswith("/image"))
    etree.SubElement(shape, tag("v", "imagedata"), {tag("r", "id"): image_rel_id})
    members = packed(doc)
    rels = etree.fromstring(members["word/_rels/document.xml.rels"])
    for identifier, kind, target, external in (("rIdFixtureChart", "chart", "charts/chart1.xml", False),
                                                ("rIdFixtureSmartArt", "diagramData", "diagrams/data1.xml", False),
                                                ("rIdFixtureExternal", "image", "https://example.invalid/fixture.png", True)):
        attributes = {"Id": identifier, "Type": NS["r"] + "/" + kind, "Target": target}
        if external:
            attributes["TargetMode"] = "External"
        etree.SubElement(rels, tag("pr", "Relationship"), attributes)
    members["word/_rels/document.xml.rels"] = xml_bytes(rels)
    chart = etree.Element(tag("c", "chartSpace"), nsmap={"c": NS["c"]})
    etree.SubElement(etree.SubElement(chart, tag("c", "chart")), tag("c", "plotArea"))
    add_member(members, "word/charts/chart1.xml", xml_bytes(chart), "application/vnd.openxmlformats-officedocument.drawingml.chart+xml")
    diagram = etree.Element(tag("dgm", "dataModel"), nsmap={"dgm": NS["dgm"]})
    etree.SubElement(diagram, tag("dgm", "ptLst"))
    etree.SubElement(diagram, tag("dgm", "cxnLst"))
    add_member(members, "word/diagrams/data1.xml", xml_bytes(diagram), "application/vnd.openxmlformats-officedocument.drawingml.diagramData+xml")
    save(members, "exclusions.docx")


if __name__ == "__main__":
    make_bench()
    make_exclusions()
    print(json.dumps({"fixtures": ["bench.docx", "bench-changed.docx", "exclusions.docx"], "directory": str(OUT)}))
